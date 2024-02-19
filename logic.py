from openai import OpenAI
import datetime
from docx import Document
todays_date = datetime.date.today()
wochentag = todays_date.strftime("%A")
#--------------------------------------------------------------------------- transcibe Audio to Text
client = OpenAI(api_key="#Insert OpenAIKey here")

def transribe_audio_to_text(audo_file_path: str) -> str:
    audio_file = open(audo_file_path, "rb")
    transcript = client.audio.transcriptions.create(
        model="whisper-1",
        file=audio_file,
        language="de"
    )
    print ("OpenAi: ")
    print (transcript.text)
    return transcript.text


#--------------------------------------------------------------------------- summarize the transcribted Text
def summarize_transcripton(text: str) -> str:
    response = client.chat.completions.create(
        model="gpt-3.5-turbo",
        messages=[
            {"role": "system", "content": "Du bist ein deutscher bot der dabei hilft ein Text von einem Programmierer in der Ausbildung in ein Tagesbericht zusammen zu fassen. schreibe aus der Ich-Perspektive"},
            {"role": "user", "content": f"Fasse mir diesen Text zusammen. Es ist ein Text welcher meinen heutigen Arbeitstag beschreibt. Mache pro punkt ein bulletpoint: {text}"}
        ]
    )
    print("OpenAi: summarized transcripted Text")
    return response.choices[0].message.content



def write_text_in_File(summarized_text: str):
    with open("/home/klaas/Downloads/Berichtsheft.txt", "a") as file:
        file.write("\n")
        file.write(format(todays_date))
        file.write(format(wochentag))
        file.write("\n")
        file.write(summarized_text)
        file.write("\n")
        file.write("--------------------------------------------------------------------------------------------------------")


def write_text_in_docx(text, wochentag):
    doc = Document('berichtsheft_template.docx')
    table = doc.tables[0]
    if wochentag == 'Montag': #Montag
        table.cell(1,1).text = text
        doc.save('berichtsheft_template.docx')
        print('finish writing in file')
    elif wochentag == 'Dienstag': #Dienstag
        table.cell(2,1).text = text
        doc.save('berichtsheft_template.docx')
        print('finish writing in file')
    elif wochentag == 'Mittwoch': #Mittwoch
        table.cell(3, 1).text = text
        doc.save('berichtsheft_template.docx')
        print('finish writing in file')
    elif wochentag == 'Donnerstag': #Donnerstag
        table.cell(4, 1).text = text
        doc.save('berichtsheft_template.docx')
        print('finish writing in file')
    elif wochentag == 'Freitag': #Freitag
        table.cell(5, 1).text = text
        doc.save('berichtsheft_template.docx')
        print('finish writing in file')
    else:
        print("error: didn't write in file ")

